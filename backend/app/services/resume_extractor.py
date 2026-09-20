"""
Resume Text Extraction Service
Supports genuine PDF (via PyMuPDF) and DOCX (via python-docx) extraction.
AI Career Preparation Agent — Academic IDP Project
Student: Chandolu Praneeth Kumar (241FA18483) — Vignan University
"""

import io
import re
from abc import ABC, abstractmethod
from typing import Optional, Tuple

try:
    import pymupdf as fitz
    PYMUPDF_AVAILABLE = True
except ImportError:
    try:
        import fitz
        PYMUPDF_AVAILABLE = True
    except ImportError:
        PYMUPDF_AVAILABLE = False

try:
    import docx  # python-docx
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit


class ResumeExtractionError(Exception):
    """Raised when resume text extraction fails due to format, corruption, or scanned images."""
    def __init__(self, message: str, code: str = "EXTRACTION_FAILED"):
        super().__init__(message)
        self.message = message
        self.code = code


def normalize_extracted_text(text: str) -> str:
    """
    Cleans raw document text:
    - Normalizes unicode quotation marks, dashes, and bullet symbols
    - Strips non-printable control characters
    - Collapses duplicate horizontal spaces
    - Collapses excessive blank lines while preserving logical paragraphs
    """
    if not text:
        return ""

    # Replace common unicode typographic characters with ASCII equivalents
    replacements = {
        "\xa0": " ",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "\n• ",
        "\u25cf": "\n• ",
        "\u25aa": "\n• ",
        "\t": "    "
    }
    for old, new in replacements.items():
        text = text.replace(old, new)

    # Strip null bytes and non-printable control characters (except newline, cr, tab)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # Collapse multiple horizontal whitespace on each line
    lines = [re.sub(r"[ ]{2,}", " ", line).strip() for line in text.splitlines()]

    # Collapse repeated empty lines (maximum 1 blank line between paragraphs)
    cleaned_lines = []
    blank_count = 0
    for line in lines:
        if not line:
            blank_count += 1
            if blank_count <= 1 and cleaned_lines:
                cleaned_lines.append("")
        else:
            blank_count = 0
            cleaned_lines.append(line)

    return "\n".join(cleaned_lines).strip()


def validate_file_bytes(file_bytes: bytes, filename: str) -> str:
    """
    Validates file size, non-emptiness, and magic bytes.
    Returns detected format ('pdf', 'docx', 'txt').
    """
    if not file_bytes or len(file_bytes) == 0:
        raise ResumeExtractionError("The uploaded file is empty (0 bytes). Please select a valid document.", "EMPTY_FILE")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise ResumeExtractionError(
            f"File size exceeds the 10 MB limit ({len(file_bytes) / (1024 * 1024):.1f} MB). Please upload a smaller file.",
            "FILE_TOO_LARGE"
        )

    # Inspect magic bytes and extension
    ext = (filename.split(".")[-1] if "." in filename else "").lower()

    # PDF detection (%PDF-)
    if file_bytes.startswith(b"%PDF"):
        return "pdf"

    # DOCX detection (ZIP header PK\x03\x04 and docx extension or [Content_Types].xml)
    if file_bytes.startswith(b"PK\x03\x04"):
        if ext in {"docx", "doc"}:
            return "docx"
        # Look for word/ or [Content_Types].xml inside first 2KB
        if b"word/" in file_bytes[:2048] or b"[Content_Types].xml" in file_bytes[:2048]:
            return "docx"

    # Plain text detection (.txt, .md, or valid UTF-8 text with common ascii chars)
    if ext in {"txt", "md"}:
        return "txt"

    # If extension matches PDF but magic header is missing -> corrupted
    if ext == "pdf":
        raise ResumeExtractionError("The PDF file header is corrupted or unreadable.", "CORRUPT_PDF")

    if ext in {"docx", "doc"}:
        raise ResumeExtractionError("The DOCX file structure is corrupted or invalid.", "CORRUPT_DOCX")

    # Reject unsupported binary formats explicitly
    unsupported_map = {
        "exe": "Executable files (.exe) are not supported.",
        "zip": "ZIP archives are not supported directly. Please upload your .pdf or .docx file.",
        "rar": "RAR archives are not supported.",
        "tar": "TAR archives are not supported.",
        "png": "Image files (.png) are not supported. Please upload a text PDF or Word document.",
        "jpg": "Image files (.jpg) are not supported. Please upload a text PDF or Word document.",
        "jpeg": "Image files (.jpeg) are not supported. Please upload a text PDF or Word document.",
        "py": "Source code files (.py) are not supported as resumes.",
        "json": "JSON files are not supported as resumes."
    }
    if ext in unsupported_map:
        raise ResumeExtractionError(unsupported_map[ext], "UNSUPPORTED_FORMAT")

    # Fallback check if it's plain text
    try:
        sample = file_bytes[:1024].decode("utf-8")
        if sample and all(ord(c) >= 32 or c in "\n\r\t" for c in sample):
            return "txt"
    except Exception:
        pass

    raise ResumeExtractionError(
        f"Unsupported file format (.{ext if ext else 'unknown'}). Please upload a PDF (.pdf) or Word document (.docx).",
        "UNSUPPORTED_FORMAT"
    )


class BaseResumeExtractor(ABC):
    """Abstract interface for resume document text extraction."""
    @abstractmethod
    def extract_text(self, file_bytes: bytes, filename: str) -> Tuple[str, str]:
        """Extracts text from file bytes. Returns (normalized_text, detected_format)."""
        pass


class DocumentResumeExtractor(BaseResumeExtractor):
    """
    Production document extractor utilizing PyMuPDF for PDF and python-docx for Word files.
    """

    def extract_text(self, file_bytes: bytes, filename: str) -> Tuple[str, str]:
        file_format = validate_file_bytes(file_bytes, filename)

        if file_format == "pdf":
            text = self._extract_pdf(file_bytes)
        elif file_format == "docx":
            text = self._extract_docx(file_bytes)
        elif file_format == "txt":
            text = self._extract_txt(file_bytes)
        else:
            raise ResumeExtractionError(f"Unsupported format: {file_format}", "UNSUPPORTED_FORMAT")

        cleaned = normalize_extracted_text(text)

        # Scanned PDF or empty extraction guard:
        # If extraction produced less than 20 characters of actual alphanumeric text
        alphanumeric_count = len(re.findall(r"[A-Za-z0-9]", cleaned))
        if alphanumeric_count < 20:
            if file_format == "pdf":
                raise ResumeExtractionError(
                    "Uploaded PDF contains no extractable text or is a scanned image without a selectable text layer. "
                    "Please upload a text-based digital PDF or DOCX file.",
                    "SCANNED_OR_EMPTY_PDF"
                )
            raise ResumeExtractionError(
                "The uploaded document contains no readable text. Please check the file contents.",
                "EMPTY_DOCUMENT_TEXT"
            )

        return cleaned, file_format

    def _extract_pdf(self, file_bytes: bytes) -> str:
        if not PYMUPDF_AVAILABLE:
            raise ResumeExtractionError("PyMuPDF engine is not available on this server.", "ENGINE_MISSING")

        try:
            stream = io.BytesIO(file_bytes)
            doc = fitz.open(stream=stream, filetype="pdf")
        except Exception as exc:
            raise ResumeExtractionError(f"Failed to open PDF document: {str(exc)}", "CORRUPT_PDF")

        pages_text = []
        try:
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                page_text = page.get_text("text") or ""
                if page_text.strip():
                    pages_text.append(page_text)
        except Exception as exc:
            raise ResumeExtractionError(f"Error reading PDF page contents: {str(exc)}", "PDF_READ_ERROR")
        finally:
            doc.close()

        return "\n\n".join(pages_text)

    def _extract_docx(self, file_bytes: bytes) -> str:
        if not PYTHON_DOCX_AVAILABLE:
            raise ResumeExtractionError("python-docx engine is not available on this server.", "ENGINE_MISSING")

        try:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
        except Exception as exc:
            raise ResumeExtractionError(f"Failed to parse DOCX document: {str(exc)}", "CORRUPT_DOCX")

        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text)

        # Also extract text from tables (many resumes format experience/skills in tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))

        return "\n".join(paragraphs)

    def _extract_txt(self, file_bytes: bytes) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore")


# Global extractor instance
resume_extractor = DocumentResumeExtractor()


def extract_resume_text(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """Convenience functional interface for resume text extraction."""
    return resume_extractor.extract_text(file_bytes, filename)
